"""Tarjetas de análisis de precio unitario con clave GuBIMclass.

Cada tarjeta se describe en tarjetas/<clave>.json (lo llena el agente de
tarjetas o una persona) y este módulo la calcula y la entrega en Excel (con
fórmulas, editable) y en PDF.

Uso: ver USO al final del archivo (base de datos, --validar, --pdf, --pendientes).

Datos que usa:
    datos/insumos.csv             precios de insumos (clave, unidad, precio, fuente)
    datos/basicos.json            básicos (morteros, concretos) desglosados en insumos
    datos/parametros_tarjeta.json indirectos, financiamiento, utilidad, salarios
    salida/construbase_gubim.csv  catálogo Construbase -> GuBIM (generar_tabla.py)
    salida/pares_tabulador_cdmx.csv  pares con el tabulador CDMX (opcional)
    fuentes/GuBIMclass_v1.2_ES.txt   nombres y jerarquía GuBIM
"""
import csv
import html
import json
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
DATOS = RAIZ / "datos"
TARJETAS = RAIZ / "tarjetas"
SALIDA = RAIZ / "salida" / "tarjetas"
PLANTILLA_CSS = RAIZ / "plantillas" / "tarjeta.css"


# --- Datos de referencia -----------------------------------------------------

def cargar_referencias():
    # insumos.csv y basicos.json más los archivos por lote (insumos_<lote>.csv,
    # basicos_<lote>.json) que escriben los agentes cuando trabajan en paralelo.
    insumos = {}
    for ruta in [DATOS / "insumos.csv"] + sorted(DATOS.glob("insumos_*.csv")):
        with open(ruta, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                r["precio"] = float(r["precio"])
                insumos.setdefault(r["clave"], r)
    basicos = {}
    for ruta in [DATOS / "basicos.json"] + sorted(DATOS.glob("basicos_*.json")):
        for k, v in json.loads(ruta.read_text(encoding="utf-8")).items():
            basicos.setdefault(k, v)
    parametros = json.loads((DATOS / "parametros_tarjeta.json").read_text(encoding="utf-8"))
    catalogo = {}
    ruta_cat = RAIZ / "salida" / "construbase_gubim.csv"
    if ruta_cat.exists():
        with open(ruta_cat, encoding="utf-8-sig") as f:
            catalogo = {r["clave_cb"]: r for r in csv.DictReader(f)}
    pares = {}
    ruta_pares = RAIZ / "salida" / "pares_tabulador_cdmx.csv"
    if ruta_pares.exists():
        with open(ruta_pares, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                # la canasta curada tiene prioridad sobre los pares automáticos
                if r["clave_cb"] not in pares or r["origen"].startswith("canasta"):
                    pares[r["clave_cb"]] = r
    gubim = {}
    ruta_g = RAIZ / "fuentes" / "GuBIMclass_v1.2_ES.txt"
    if ruta_g.exists():
        for linea in ruta_g.read_text(encoding="utf-8").splitlines():
            p = linea.split("\t")
            if len(p) >= 3:
                gubim[p[0]] = p[1]
    return {"insumos": insumos, "basicos": basicos, "parametros": parametros,
            "catalogo": catalogo, "pares": pares, "gubim": gubim}


# --- Cálculo ---------------------------------------------------------------------

def r2(x):
    """Redondeo a centavos igual al de Excel (mitad hacia arriba)."""
    from decimal import Decimal, ROUND_HALF_UP
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


class ErrorTarjeta(Exception):
    pass


def calcular(t, ref):
    """Devuelve la tarjeta calculada (dict) a partir de su definición JSON."""
    ins, bas, par = ref["insumos"], ref["basicos"], ref["parametros"]
    cat = ref["catalogo"].get(t["clave_cb"], {})
    faltan = []

    materiales = []
    for m in t.get("materiales", []):
        cant = m["cantidad"] * (1 + m.get("desperdicio", 0))
        if "basico" in m:
            b = bas.get(m["basico"])
            if not b:
                faltan.append("básico " + m["basico"])
                continue
            comps = []
            for clave, coef in b["componentes"]:
                i = ins.get(clave)
                if not i:
                    faltan.append("insumo " + clave)
                    continue
                q = coef * cant
                comps.append({"clave": clave, "descripcion": i["descripcion"], "unidad": i["unidad"],
                              "coeficiente": coef, "cantidad": q, "precio": i["precio"],
                              "importe": r2(q * i["precio"]), "fuente": i["fuente"], "estado": i["estado"]})
            materiales.append({"tipo": "basico", "clave": m["basico"], "descripcion": b["descripcion"],
                               "unidad": b["unidad"], "cantidad_base": m["cantidad"],
                               "desperdicio": m.get("desperdicio", 0), "cantidad": cant,
                               "componentes": comps, "importe": r2(sum(c["importe"] for c in comps)),
                               "nota": m.get("nota", "")})
        else:
            i = ins.get(m["insumo"])
            if not i:
                faltan.append("insumo " + m["insumo"])
                continue
            materiales.append({"tipo": "insumo", "clave": m["insumo"], "descripcion": i["descripcion"],
                               "unidad": i["unidad"], "cantidad_base": m["cantidad"],
                               "desperdicio": m.get("desperdicio", 0), "cantidad": cant,
                               "precio": i["precio"], "importe": r2(cant * i["precio"]),
                               "nota": m.get("nota", ""), "fuente": i["fuente"], "estado": i["estado"]})
    if faltan:
        raise ErrorTarjeta("Faltan datos de referencia: " + ", ".join(sorted(set(faltan))))
    sub_mat = r2(sum(m["importe"] for m in materiales))

    mo = t["mano_obra"]
    cuadrilla = []
    for c in mo["cuadrilla"]:
        nombre, salario = par["categorias"][c["categoria"]]
        cuadrilla.append({"categoria": c["categoria"], "nombre": nombre, "personas": c["personas"],
                          "salario": salario, "importe_jornada": r2(c["personas"] * salario)})
    costo_cuadrilla = r2(sum(c["importe_jornada"] for c in cuadrilla))
    rend = mo["rendimiento"]
    sub_mo = r2(costo_cuadrilla / rend)

    equipo = [{"descripcion": "Herramienta menor", "base": "% de mano de obra",
               "factor": par["herramienta_menor"], "importe": r2(sub_mo * par["herramienta_menor"])}]
    for e in t.get("equipo", []):
        if "horas" in e:   # equipo por costo horario: horas por unidad × costo horario
            equipo.append({"descripcion": e["descripcion"], "base": f"{e['horas']:.4f} h × costo horario",
                           "horas": e["horas"], "costo_horario": e["costo_horario"],
                           "importe": r2(e["horas"] * e["costo_horario"])})
    sub_eq = r2(sum(e["importe"] for e in equipo))

    cd = r2(sub_mat + sub_mo + sub_eq)
    ind = r2(cd * par["indirectos"])
    fin = r2((cd + ind) * par["financiamiento"])
    uti = r2((cd + ind + fin) * par["utilidad"])
    car = r2((cd + ind + fin + uti) * par["cargos_adicionales"])
    pu = r2(cd + ind + fin + uti + car)
    personas = sum(c["personas"] for c in cuadrilla)

    ruta = []
    codigo = t["gubim"]
    partes = codigo.split(".")
    for n in range(1, len(partes)):
        k = ".".join(partes[:n])
        ruta.append((k, ref["gubim"].get(k, "")))

    par_cdmx = ref["pares"].get(t["clave_cb"])
    return {
        "def": t, "cat": cat, "materiales": materiales, "sub_mat": sub_mat,
        "cuadrilla": cuadrilla, "costo_cuadrilla": costo_cuadrilla, "rendimiento": rend,
        "fuente_rend": mo.get("fuente", ""), "sub_mo": sub_mo, "equipo": equipo, "sub_eq": sub_eq,
        "cd": cd, "ind": ind, "fin": fin, "uti": uti, "car": car, "pu": pu, "par": par,
        "jornadas_u": 1 / rend, "hh_u": personas * 8 / rend, "jornadas_100": 100 / rend,
        "gubim_nombre": ref["gubim"].get(codigo, ""), "gubim_ruta": ruta,
        "comparativo": {
            "pu_2017": float(cat["pu_2017"]) if cat.get("pu_2017") else None,
            "pu_act": float(cat["pu_actualizado"]) if cat.get("pu_actualizado") else None,
            "pu_cdmx": float(par_cdmx["pu_tabulador"]) if par_cdmx else None,
            "clave_cdmx": par_cdmx["clave_tabulador"] if par_cdmx else None,
        },
    }


# --- Importe con letra --------------------------------------------------------------

_UNI = ["", "UN", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE", "DIEZ",
        "ONCE", "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISÉIS", "DIECISIETE", "DIECIOCHO",
        "DIECINUEVE", "VEINTE", "VEINTIÚN", "VEINTIDÓS", "VEINTITRÉS", "VEINTICUATRO", "VEINTICINCO",
        "VEINTISÉIS", "VEINTISIETE", "VEINTIOCHO", "VEINTINUEVE"]
_DEC = ["", "", "", "TREINTA", "CUARENTA", "CINCUENTA", "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
_CEN = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS", "SEISCIENTOS",
        "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]


def _cientos(n):
    if n == 100:
        return "CIEN"
    c, r = divmod(n, 100)
    partes = [_CEN[c]] if c else []
    if r < 30:
        partes.append(_UNI[r])
    else:
        d, u = divmod(r, 10)
        partes.append(_DEC[d] + (" Y " + _UNI[u] if u else ""))
    return " ".join(p for p in partes if p)


def numero_a_letras(n):
    n = int(n)
    if n == 0:
        return "CERO"
    millones, resto = divmod(n, 1_000_000)
    miles, unidades = divmod(resto, 1000)
    partes = []
    if millones:
        partes.append("UN MILLÓN" if millones == 1 else _cientos(millones) + " MILLONES")
    if miles:
        partes.append("MIL" if miles == 1 else _cientos(miles) + " MIL")
    if unidades:
        partes.append(_cientos(unidades))
    return " ".join(partes)


def importe_con_letra(x):
    pesos = int(x)
    centavos = int(round((x - pesos) * 100))
    texto = numero_a_letras(pesos)
    moneda = "PESO" if pesos == 1 else "PESOS"
    if texto.endswith("MILLÓN") or texto.endswith("MILLONES"):
        moneda = "DE " + moneda
    return f"({texto} {moneda} {centavos:02d}/100 M.N.)"


# --- Formato -----------------------------------------------------------------------

def mx(x, dec=2):
    return f"{x:,.{dec}f}"


def pct(x):
    return f"{x * 100:.2f} %"


UNIDADES_VISTA = {"M2": "m²", "M3": "m³", "M": "m", "PZA": "pza", "KG": "kg", "TON": "ton", "SAL": "salida",
                  "JGO": "juego", "M3/KM": "m³-km", "HA": "ha", "LOTE": "lote"}


def unidad_vista(u):
    return UNIDADES_VISTA.get(u, u.lower() if u else "")


# --- HTML / PDF ------------------------------------------------------------------------

FUENTES = [("IBM Plex Sans", 400, "IBMPlexSans-Regular"), ("IBM Plex Sans", 500, "IBMPlexSans-Medium"),
           ("IBM Plex Sans", 600, "IBMPlexSans-SemiBold"), ("IBM Plex Mono", 400, "IBMPlexMono-Regular"),
           ("IBM Plex Mono", 500, "IBMPlexMono-Medium"), ("IBM Plex Sans Condensed", 500, "IBMPlexSansCondensed-Medium"),
           ("IBM Plex Sans Condensed", 600, "IBMPlexSansCondensed-SemiBold")]


def fuentes_css():
    """@font-face con las fuentes incrustadas (IBM Plex, licencia OFL), para que el
    PDF salga igual sin depender de internet."""
    import base64
    carpeta = RAIZ / "plantillas" / "fuentes"
    reglas = []
    for familia, peso, archivo in FUENTES:
        ruta = carpeta / f"{archivo}.woff2"
        if ruta.exists():
            datos = base64.b64encode(ruta.read_bytes()).decode()
            reglas.append(f"@font-face{{font-family:'{familia}';font-weight:{peso};font-style:normal;"
                          f"src:url(data:font/woff2;base64,{datos}) format('woff2')}}")
    return "".join(reglas)


def html_tarjeta(c):
    t, cat, par = c["def"], c["cat"], c["par"]
    e = html.escape
    u = unidad_vista(cat.get("unidad", ""))
    filas_mat = []
    for m in c["materiales"]:
        desp = f" + {m['desperdicio'] * 100:.0f} % desperdicio" if m["desperdicio"] else ""
        nota = e(m["nota"] + desp)
        if m["tipo"] == "insumo":
            filas_mat.append(
                f"<tr><td>{e(m['descripcion'])}<span class='src'>{nota}</span></td><td>{e(m['unidad'])}</td>"
                f"<td class='n'>{mx(m['cantidad'], 4)}</td><td class='n'>{mx(m['precio'])}</td>"
                f"<td class='n'>{mx(m['importe'])}</td></tr>")
        else:
            filas_mat.append(
                f"<tr class='basico'><td>{e(m['descripcion'])} <span class='tag'>básico</span>"
                f"<span class='src'>{nota}</span></td><td>{e(m['unidad'])}</td>"
                f"<td class='n'>{mx(m['cantidad'], 4)}</td><td class='n'></td><td class='n'>{mx(m['importe'])}</td></tr>")
            for k in m["componentes"]:
                filas_mat.append(
                    f"<tr class='comp'><td>{e(k['descripcion'])}<span class='src'>{mx(k['coeficiente'], 3)} "
                    f"{e(k['unidad'])} por {e(m['unidad'])} de {e(m['descripcion'].lower())}</span></td>"
                    f"<td>{e(k['unidad'])}</td><td class='n'>{mx(k['cantidad'], 4)}</td>"
                    f"<td class='n'>{mx(k['precio'])}</td><td class='n'>{mx(k['importe'])}</td></tr>")
    n = len(c["cuadrilla"])
    filas_mo = []
    for i, q in enumerate(c["cuadrilla"]):
        extra = (f"<td class='n' rowspan='{n}'>{mx(c['costo_cuadrilla'])}</td>"
                 f"<td class='n' rowspan='{n}'>{mx(c['rendimiento'])} {u}/jor</td>"
                 f"<td class='n' rowspan='{n}'>{mx(c['sub_mo'])}</td>") if i == 0 else ""
        src = f"<span class='src'>{e(par['nota_salarios'])}</span>" if i == 0 else ""
        filas_mo.append(f"<tr><td>{e(q['nombre'])}{src}</td><td class='n'>{q['personas']}</td>"
                        f"<td class='n'>{mx(q['salario'])}</td>{extra}</tr>")
    filas_eq = "".join(
        f"<tr><td>{e(q['descripcion'])}</td><td>{e(q['base'])}</td>"
        f"<td class='n'>{pct(q['factor']) if 'factor' in q else mx(q['costo_horario'])}</td>"
        f"<td class='n'>{mx(q['importe'])}</td></tr>" for q in c["equipo"])
    ruta = "".join(f"<li><span>{e(k)}</span>{e(v)}</li>" for k, v in c["gubim_ruta"])
    alts = "; ".join(f"<b>{e(a['clave'])}</b> {e(a['condicion'])}" for a in t.get("gubim_alternativas", []))
    cd = c["cd"]
    comp = [(c["sub_mat"] / cd, "var(--mat)", "Materiales"), (c["sub_mo"] / cd, "var(--mo)", "Mano de obra"),
            (c["sub_eq"] / cd, "var(--eq)", "Herramienta y equipo")]
    barra = "".join(f"<i style='width:{p * 100:.1f}%;background:{col}'></i>" for p, col, _ in comp)
    leyenda = "".join(f"<span><i class='sw' style='background:{col}'></i>{nom} {p * 100:.1f} %</span>"
                      for p, col, nom in comp)
    cmp_rows = []
    k = c["comparativo"]
    for etiqueta, valor in (("Construbase 2017", k["pu_2017"]),
                            ("Construbase actualizado " + par["fecha_base"], k["pu_act"]),
                            (f"Tabulador CDMX jul-2026 · {k['clave_cdmx']}" if k["clave_cdmx"] else None, k["pu_cdmx"])):
        if etiqueta and valor:
            delta = f"{(c['pu'] / valor - 1) * 100:+.1f} %"
            cmp_rows.append(f"<tr><td>{e(etiqueta)}</td><td class='n'>{mx(valor)}</td><td class='n delta'>{delta}</td></tr>")
    cmp_rows.append(f"<tr class='me'><td>Esta tarjeta</td><td class='n'>{mx(c['pu'])}</td><td></td></tr>")
    supuestos = "".join(f"<li>{e(s)}</li>" for s in t.get("supuestos", []))
    estado = "Revisada" if t.get("reviso") else "Borrador del agente · pendiente de revisión"
    capitulo = " › ".join(x for x in (cat.get("capitulo", "").title(), cat.get("subcapitulo", "").title()) if x)
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<title>Tarjeta {e(t['clave_cb'])}</title>
<style>{fuentes_css()}{PLANTILLA_CSS.read_text(encoding='utf-8')}</style></head><body>
<article class="card">
<header class="head"><div><div class="firm">{e(par['empresa'])}</div>
<div class="doc-title">Tarjeta de análisis de precio unitario</div><span class="status">{e(estado)}</span></div>
<dl class="meta"><dt>Tarjeta</dt><dd>TPU-{e(t['clave_cb'])}</dd><dt>Fecha base</dt><dd>{e(par['fecha_base'])}</dd>
<dt>Zona</dt><dd>{e(par['zona'])}</dd><dt>Proyecto</dt><dd>{e(t.get('proyecto', 'Catálogo base'))}</dd></dl></header>
<div class="ident"><div class="concept"><div class="label">Concepto</div>
<div class="keyrow"><span class="key">{e(t['clave_cb'])}</span><span class="unit">{e(u)}</span><span class="label">{e(capitulo)}</span></div>
<p>{e(cat.get('descripcion', ''))}</p></div>
<aside class="gubim"><div class="label">GuBIMclass v1.2 · Assembly Code</div>
<div class="gcode">{e(t['gubim'])}</div><div class="gname">{e(c['gubim_nombre'])}</div>
<ul class="path">{ruta}</ul>
<div class="chips"><span class="chip">Nivel {len(t['gubim'].split('.'))}</span><span class="chip">Confianza {e(t.get('gubim_confianza', ''))}</span></div>
{f'<div class="alt">Alternativas: {alts}.</div>' if alts else ''}</aside></div>
<section class="block"><div class="bh"><h2><i class="sw" style="background:var(--mat)"></i>A. Materiales</h2><span class="label">Cantidades con desperdicio · básicos desglosados</span></div>
<table><thead><tr><th>Insumo</th><th>Unidad</th><th class="n">Cantidad</th><th class="n">Precio unitario</th><th class="n">Importe</th></tr></thead>
<tbody>{''.join(filas_mat)}<tr class="sub"><td colspan="4">Subtotal materiales</td><td class="n">{mx(c['sub_mat'])}</td></tr></tbody></table></section>
<section class="block"><div class="bh"><h2><i class="sw" style="background:var(--mo)"></i>B. Mano de obra</h2><span class="formula">importe = costo cuadrilla ÷ rendimiento</span></div>
<table><thead><tr><th>Categoría</th><th class="n">Personas</th><th class="n">Salario real / jornada</th><th class="n">Costo cuadrilla</th><th class="n">Rendimiento</th><th class="n">Importe</th></tr></thead>
<tbody>{''.join(filas_mo)}<tr class="sub"><td colspan="5">Subtotal mano de obra <span class="src inline">{e(c['fuente_rend'])}</span></td><td class="n">{mx(c['sub_mo'])}</td></tr></tbody></table></section>
<section class="block"><div class="bh"><h2><i class="sw" style="background:var(--eq)"></i>C. Herramienta y equipo</h2></div>
<table><thead><tr><th>Concepto</th><th>Base</th><th class="n">Factor</th><th class="n">Importe</th></tr></thead>
<tbody>{filas_eq}<tr class="sub"><td colspan="3">Subtotal herramienta y equipo</td><td class="n">{mx(c['sub_eq'])}</td></tr></tbody></table></section>
<div class="bottom"><div class="summary"><div class="label">Integración del precio</div><table><tbody>
<tr><td>Materiales (A)</td><td class="n">{mx(c['sub_mat'])}</td></tr><tr><td>Mano de obra (B)</td><td class="n">{mx(c['sub_mo'])}</td></tr>
<tr><td>Herramienta y equipo (C)</td><td class="n">{mx(c['sub_eq'])}</td></tr><tr class="cd"><td>Costo directo</td><td class="n">{mx(cd)}</td></tr>
<tr><td>Indirectos {pct(par['indirectos'])}</td><td class="n">{mx(c['ind'])}</td></tr><tr><td>Financiamiento {pct(par['financiamiento'])}</td><td class="n">{mx(c['fin'])}</td></tr>
<tr><td>Utilidad {pct(par['utilidad'])}</td><td class="n">{mx(c['uti'])}</td></tr><tr><td>Cargos adicionales {pct(par['cargos_adicionales'])}</td><td class="n">{mx(c['car'])}</td></tr>
</tbody></table><div class="pu"><span class="label">Precio unitario / {e(u)}</span><span class="big">${mx(c['pu'])}</span></div>
<div class="words">{importe_con_letra(c['pu'])} · No incluye IVA</div></div>
<div class="side"><div><div class="label">Composición del costo directo</div><div class="bar">{barra}</div><div class="legend">{leyenda}</div></div>
<div><div class="label">Tiempo de ejecución</div><div class="time">
<div><div class="v">{mx(c['jornadas_u'], 4)}</div><div class="l">jornadas de cuadrilla / {e(u)}</div></div>
<div><div class="v">{mx(c['hh_u'])}</div><div class="l">horas-hombre / {e(u)}</div></div>
<div><div class="v">{mx(c['jornadas_100'], 1)}</div><div class="l">jornadas por 100 {e(u)}</div></div></div></div>
<div class="cmp"><div class="label">Comparativo de referencia</div><table><tbody>{''.join(cmp_rows)}</tbody></table>
<div class="words">Diferencia de esta tarjeta contra cada referencia.</div></div>
{f'<div><div class="label">Supuestos</div><ul class="sup">{supuestos}</ul></div>' if supuestos else ''}</div></div>
<footer class="foot"><span>Precios de insumos: <b>referencia, por validar</b> · Rendimiento: <b>{e(c['fuente_rend'] or 'tabla de referencia')}</b></span>
<span>Elaboró: <b>{e(t.get('elaboro', ''))}</b> · Revisó: <b>{e(t.get('reviso') or '—')}</b></span></footer>
</article></body></html>"""


def pdf_desde_html(documento, destino):
    import glob
    from playwright.sync_api import sync_playwright

    candidatos = glob.glob("/opt/pw-browsers/chromium-*/chrome-linux*/chrome")
    with sync_playwright() as p:
        nav = p.chromium.launch(executable_path=candidatos[0]) if candidatos else p.chromium.launch()
        pagina = nav.new_page()
        pagina.set_content(documento, wait_until="load")
        pagina.pdf(path=str(destino), format="Letter", print_background=True, scale=0.86,
                   margin={"top": "10mm", "bottom": "10mm", "left": "10mm", "right": "10mm"})
        nav.close()


# --- Excel -------------------------------------------------------------------------------

def excel_tarjeta(c, destino):
    """Excel con la misma estructura y fórmulas: cantidades, precios, salarios,
    rendimiento y porcentajes se pueden editar y el P.U. se recalcula."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    t, cat, par = c["def"], c["cat"], c["par"]
    wb = Workbook()
    ws = wb.active
    ws.title = t["clave_cb"]
    azul, verde = "1F4E78", "0F6B5C"
    enc = PatternFill("solid", fgColor=azul)
    suave = PatternFill("solid", fgColor="E3ECF5")
    gub = PatternFill("solid", fgColor="E0F1ED")
    edit = PatternFill("solid", fgColor="FFF2CC")
    blanco = Font(bold=True, color="FFFFFF")
    fino = Side(style="thin", color="9AA7B4")
    for col, ancho in zip("ABCDEFG", (48, 10, 12, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = ancho
    fila = [0]

    def add(valores=(), bold=False, fill=None, fmt=None):
        fila[0] += 1
        r = fila[0]
        for i, v in enumerate(valores, start=1):
            celda = ws.cell(row=r, column=i, value=v)
            if bold:
                celda.font = Font(bold=True)
            if fill:
                celda.fill = fill
        if fmt:
            for col, f in fmt.items():
                ws.cell(row=r, column=col).number_format = f
        return r

    def encabezado(titulos):
        r = add(titulos)
        for i in range(1, len(titulos) + 1):
            ws.cell(row=r, column=i).fill = enc
            ws.cell(row=r, column=i).font = blanco
        return r

    r = add([par["empresa"]], bold=True)
    ws.cell(row=r, column=1).font = Font(bold=True, color=azul, size=12)
    r = add(["Tarjeta de análisis de precio unitario"], bold=True)
    ws.cell(row=r, column=1).font = Font(bold=True, size=14)
    add(["Tarjeta", f"TPU-{t['clave_cb']}", "", "Fecha base", par["fecha_base"], "Zona", par["zona"]])
    add(["Estado", "Revisada" if t.get("reviso") else "Borrador del agente, pendiente de revisión"])
    add()
    r = add(["Concepto " + t["clave_cb"] + " · " + unidad_vista(cat.get("unidad", ""))], bold=True, fill=suave)
    r = add([cat.get("descripcion", "")])
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 45
    r = add([f"GuBIMclass {t['gubim']} · {c['gubim_nombre']}", "", "", "Confianza", t.get("gubim_confianza", "")],
            bold=True, fill=gub)
    ws.cell(row=r, column=1).font = Font(bold=True, color=verde)
    add([" › ".join(f"{k} {v}" for k, v in c["gubim_ruta"])])
    for a in t.get("gubim_alternativas", []):
        add([f"Alternativa {a['clave']}: {a['condicion']}"])
    add()

    # A. Materiales
    add(["A. Materiales"], bold=True)
    encabezado(["Insumo", "Unidad", "Cant. base", "Desperdicio", "Cantidad", "Precio unitario", "Importe"])
    filas_importe = []
    for m in c["materiales"]:
        if m["tipo"] == "insumo":
            r = add([m["descripcion"], m["unidad"], m["cantidad_base"], m["desperdicio"]])
            ws.cell(row=r, column=5, value=f"=C{r}*(1+D{r})")
            ws.cell(row=r, column=6, value=m["precio"]).fill = edit
            ws.cell(row=r, column=7, value=f"=ROUND(E{r}*F{r},2)")
            for col in (3, 4):
                ws.cell(row=r, column=col).fill = edit
            filas_importe.append(r)
        else:
            rb = add([m["descripcion"] + " (básico)", m["unidad"], m["cantidad_base"], m["desperdicio"]], bold=True)
            ws.cell(row=rb, column=5, value=f"=C{rb}*(1+D{rb})")
            for col in (3, 4):
                ws.cell(row=rb, column=col).fill = edit
            comps = []
            for k in m["componentes"]:
                r = add(["    " + k["descripcion"], k["unidad"], k["coeficiente"], "coef."])
                ws.cell(row=r, column=5, value=f"=C{r}*E{rb}")
                ws.cell(row=r, column=6, value=k["precio"]).fill = edit
                ws.cell(row=r, column=7, value=f"=ROUND(E{r}*F{r},2)")
                ws.cell(row=r, column=3).fill = edit
                comps.append(r)
            ws.cell(row=rb, column=7, value=f"=SUM(G{comps[0]}:G{comps[-1]})")
            filas_importe.append(rb)
        for rr in range(r if m["tipo"] == "insumo" else rb, fila[0] + 1):
            ws.cell(row=rr, column=3).number_format = "0.0000"
            ws.cell(row=rr, column=5).number_format = "0.0000"
            ws.cell(row=rr, column=6).number_format = "#,##0.00"
            ws.cell(row=rr, column=7).number_format = "#,##0.00"
            if ws.cell(row=rr, column=4).value != "coef.":
                ws.cell(row=rr, column=4).number_format = "0%"
    r_mat = add(["Subtotal materiales"], bold=True)
    ws.cell(row=r_mat, column=7, value="=" + "+".join(f"G{x}" for x in filas_importe)).number_format = "#,##0.00"
    add()

    # B. Mano de obra
    add(["B. Mano de obra"], bold=True)
    encabezado(["Categoría", "Personas", "Salario real / jornada", "Importe jornada"])
    filas_cuad = []
    for q in c["cuadrilla"]:
        r = add([q["nombre"], q["personas"], q["salario"]])
        ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = "#,##0.00"
        ws.cell(row=r, column=3).number_format = "#,##0.00"
        ws.cell(row=r, column=2).fill = edit
        ws.cell(row=r, column=3).fill = edit
        filas_cuad.append(r)
    r_cc = add(["Costo de la cuadrilla por jornada"])
    ws.cell(row=r_cc, column=4, value=f"=SUM(D{filas_cuad[0]}:D{filas_cuad[-1]})").number_format = "#,##0.00"
    r_rend = add([f"Rendimiento ({unidad_vista(cat.get('unidad', ''))} por jornada) · {c['fuente_rend']}"])
    ws.cell(row=r_rend, column=4, value=c["rendimiento"]).fill = edit
    r_mo = add(["Subtotal mano de obra"], bold=True)
    ws.cell(row=r_mo, column=7, value=f"=ROUND(D{r_cc}/D{r_rend},2)").number_format = "#,##0.00"
    add()

    # C. Herramienta y equipo
    add(["C. Herramienta y equipo"], bold=True)
    filas_eq = []
    r = add(["Herramienta menor (% de mano de obra)", "", par["herramienta_menor"]])
    ws.cell(row=r, column=3).number_format = "0.00%"
    ws.cell(row=r, column=3).fill = edit
    ws.cell(row=r, column=7, value=f"=ROUND(G{r_mo}*C{r},2)").number_format = "#,##0.00"
    filas_eq.append(r)
    for q in c["equipo"][1:]:
        r = add([q["descripcion"] + " (horas por unidad × costo horario)", "", q["horas"], "", "", q["costo_horario"]])
        ws.cell(row=r, column=7, value=f"=ROUND(C{r}*F{r},2)").number_format = "#,##0.00"
        filas_eq.append(r)
    r_eq = add(["Subtotal herramienta y equipo"], bold=True)
    ws.cell(row=r_eq, column=7, value="=" + "+".join(f"G{x}" for x in filas_eq)).number_format = "#,##0.00"
    add()

    # Integración
    add(["Integración del precio"], bold=True)
    r_cd = add(["Costo directo"], bold=True, fill=suave)
    ws.cell(row=r_cd, column=7, value=f"=G{r_mat}+G{r_mo}+G{r_eq}")
    filas = {}
    for clave, etiqueta, base in (("indirectos", "Indirectos", f"G{r_cd}"),
                                  ("financiamiento", "Financiamiento", None),
                                  ("utilidad", "Utilidad", None),
                                  ("cargos_adicionales", "Cargos adicionales", None)):
        r = add([etiqueta, "", par[clave]])
        ws.cell(row=r, column=3).number_format = "0.00%"
        ws.cell(row=r, column=3).fill = edit
        filas[clave] = r
    ri, rf, ru, rc = filas["indirectos"], filas["financiamiento"], filas["utilidad"], filas["cargos_adicionales"]
    ws.cell(row=ri, column=7, value=f"=ROUND(G{r_cd}*C{ri},2)")
    ws.cell(row=rf, column=7, value=f"=ROUND((G{r_cd}+G{ri})*C{rf},2)")
    ws.cell(row=ru, column=7, value=f"=ROUND((G{r_cd}+G{ri}+G{rf})*C{ru},2)")
    ws.cell(row=rc, column=7, value=f"=ROUND((G{r_cd}+G{ri}+G{rf}+G{ru})*C{rc},2)")
    r_pu = add([f"PRECIO UNITARIO / {unidad_vista(cat.get('unidad', ''))} (sin IVA)"], bold=True, fill=enc)
    ws.cell(row=r_pu, column=1).font = blanco
    ws.cell(row=r_pu, column=7, value=f"=G{r_cd}+G{ri}+G{rf}+G{ru}+G{rc}")
    ws.cell(row=r_pu, column=7).font = blanco
    ws.cell(row=r_pu, column=7).fill = enc
    for rr in (r_cd, ri, rf, ru, rc, r_pu):
        ws.cell(row=rr, column=7).number_format = "$#,##0.00"
    add([importe_con_letra(c["pu"]) + "  (al valor calculado por el agente)"])
    add()
    add(["Tiempo de ejecución"], bold=True)
    personas = sum(q["personas"] for q in c["cuadrilla"])
    r = add(["Jornadas de cuadrilla por unidad"])
    ws.cell(row=r, column=7, value=f"=1/D{r_rend}").number_format = "0.0000"
    r = add(["Horas-hombre por unidad"])
    ws.cell(row=r, column=7, value=f"={personas}*8/D{r_rend}").number_format = "0.00"
    add()
    add(["Comparativo de referencia"], bold=True)
    k = c["comparativo"]
    for etiqueta, valor in (("Construbase 2017", k["pu_2017"]), ("Construbase actualizado", k["pu_act"]),
                            (f"Tabulador CDMX jul-2026 · {k['clave_cdmx']}", k["pu_cdmx"])):
        if valor:
            r = add([etiqueta, "", "", "", "", valor])
            ws.cell(row=r, column=6).number_format = "$#,##0.00"
            ws.cell(row=r, column=7, value=f"=G{r_pu}/F{r}-1").number_format = "+0.0%;-0.0%"
    if t.get("supuestos"):
        add()
        add(["Supuestos"], bold=True)
        for s in t["supuestos"]:
            add(["• " + s])
    add()
    add([f"Elaboró: {t.get('elaboro', '')}", "", "", f"Revisó: {t.get('reviso') or '—'}"])
    add(["Celdas amarillas: datos editables (cantidades, precios, salarios, rendimiento, porcentajes)."])
    for row in ws.iter_rows(min_row=1, max_row=fila[0]):
        for celda in row:
            if celda.column > 1 and celda.value is not None and not isinstance(celda.value, str):
                celda.border = Border(bottom=fino)
    ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)


# --- Base de datos -------------------------------------------------------------------

BASE_SQLITE = RAIZ / "salida" / "tarjetas.sqlite"
BASE_EXCEL = RAIZ / "salida" / "Base_tarjetas_PU.xlsx"

COLUMNAS_TARJETA = [
    "clave_cb", "gubim", "gubim_catalogo", "gubim_nombre", "gubim_confianza", "gubim_alternativas", "capitulo", "subcapitulo",
    "descripcion", "unidad", "materiales", "mano_obra", "herramienta_equipo", "costo_directo", "indirectos",
    "financiamiento", "utilidad", "cargos_adicionales", "pu", "cuadrilla", "costo_cuadrilla", "rendimiento",
    "jornadas_por_unidad", "hh_por_unidad", "rendimiento_implicito_cdmx", "pu_construbase_2017", "pu_actualizado", "pu_cdmx", "clave_cdmx",
    "dif_vs_actualizado", "dif_vs_cdmx", "dif_vs_referencia", "referencia_validacion", "referencia_sustituta", "cdmx_no_comparable", "desviacion_justificada", "alerta", "estado", "elaboro", "reviso", "fecha_base", "supuestos",
]
INDIRECTO_CDMX = 0.2751  # indirecto integrado del tabulador CDMX 2026 (sus P.U. ya sin cargos adicionales)
UMBRAL_ALERTA = 0.25   # diferencia contra la referencia que merece revisión
TOPE_JUSTIFICABLE = 0.50  # arriba de esto la alerta queda aunque la tarjeta traiga justificación


def fila_tarjeta(c):
    t, cat, k = c["def"], c["cat"], c["comparativo"]
    # Los básicos de Construbase (conceptos en mayúsculas: concretos, morteros) están
    # a costo directo, igual que los Básicos (BAS) del tabulador CDMX: esas tarjetas
    # comparan su costo directo contra ambas referencias.
    a_cd = bool(t.get("comparar_costo_directo"))
    base = c["cd"] if a_cd else c["pu"]
    dif_act = base / k["pu_act"] - 1 if k["pu_act"] else None
    dif_cdmx = base / k["pu_cdmx"] - 1 if k["pu_cdmx"] else None
    # La referencia es la CDMX cuando hay par comparable; si la tarjeta marca el
    # par como no comparable (otro alcance o sistema), se usa el P.U. actualizado.
    no_comparable = t.get("cdmx_no_comparable", "")
    usa_cdmx = dif_cdmx is not None and not no_comparable
    ref = dif_cdmx if usa_cdmx else dif_act
    nombre_ref = ("CDMX" if usa_cdmx else "P.U. actualizado") + (" (a costo directo)" if a_cd else "")
    # Si el P.U. actualizado del catálogo es erróneo (rompe la progresión de su familia,
    # incluye otro alcance), la tarjeta puede dar una referencia sustituta con su razón.
    sust = t.get("referencia_sustituta")
    if sust and not usa_cdmx:
        ref, nombre_ref = c["pu"] / sust["pu"] - 1, "Referencia sustituta"
    # Rendimiento que implica el precio de la CDMX: su costo directo menos los
    # materiales y equipo de esta tarjeta deja la mano de obra (con herramienta).
    rend_cdmx = None
    if k["pu_cdmx"]:
        cd_cdmx = k["pu_cdmx"] if a_cd else k["pu_cdmx"] / (1 + INDIRECTO_CDMX)
        otros = sum(e["importe"] for e in c["equipo"][1:])
        mo_cdmx = (cd_cdmx - c["sub_mat"] - otros) / (1 + c["par"]["herramienta_menor"])
        # Solo tiene sentido si el precio CDMX deja una mano de obra razonable
        # (entre 1/4 y 4 veces el rendimiento de la tarjeta); si no, el par no es comparable.
        if mo_cdmx > 0 and c["rendimiento"] / 4 <= c["costo_cuadrilla"] / mo_cdmx <= c["rendimiento"] * 4:
            rend_cdmx = round(c["costo_cuadrilla"] / mo_cdmx, 2)
    # Una diferencia explicada (precio regional, otra especificación) se documenta en
    # "desviacion_justificada" y deja de contar como alerta, hasta TOPE_JUSTIFICABLE.
    justificada = t.get("desviacion_justificada", "")
    alerta = ""
    if ref is not None and abs(ref) > UMBRAL_ALERTA:
        if not justificada or abs(ref) > TOPE_JUSTIFICABLE:
            alerta = "Revisar: {:+.0%} contra {}".format(ref, nombre_ref)
    else:
        justificada = ""
    return {
        "clave_cb": t["clave_cb"], "gubim": t["gubim"], "gubim_catalogo": cat.get("clave_gubim", ""),
        "gubim_nombre": c["gubim_nombre"],
        "gubim_confianza": t.get("gubim_confianza", ""),
        "gubim_alternativas": "; ".join(f"{a['clave']} {a['condicion']}" for a in t.get("gubim_alternativas", [])),
        "capitulo": cat.get("capitulo", ""), "subcapitulo": cat.get("subcapitulo", ""),
        "descripcion": cat.get("descripcion", ""), "unidad": cat.get("unidad", ""),
        "materiales": c["sub_mat"], "mano_obra": c["sub_mo"], "herramienta_equipo": c["sub_eq"],
        "costo_directo": c["cd"], "indirectos": c["ind"], "financiamiento": c["fin"], "utilidad": c["uti"],
        "cargos_adicionales": c["car"], "pu": c["pu"],
        "cuadrilla": " + ".join(f"{q['personas']} {q['nombre']}" for q in c["cuadrilla"]),
        "costo_cuadrilla": c["costo_cuadrilla"], "rendimiento": c["rendimiento"],
        "jornadas_por_unidad": round(c["jornadas_u"], 6), "hh_por_unidad": round(c["hh_u"], 4),
        "pu_construbase_2017": k["pu_2017"], "pu_actualizado": k["pu_act"], "pu_cdmx": k["pu_cdmx"],
        "clave_cdmx": k["clave_cdmx"] or "",
        "rendimiento_implicito_cdmx": rend_cdmx,
        "dif_vs_actualizado": round(dif_act, 4) if dif_act is not None else None,
        "dif_vs_cdmx": round(dif_cdmx, 4) if dif_cdmx is not None else None,
        "dif_vs_referencia": round(ref, 4) if ref is not None else None,
        "referencia_validacion": nombre_ref if ref is not None else "",
        "referencia_sustituta": f"{sust['pu']:,.2f}: {sust['razon']}" if sust and not usa_cdmx else "",
        "cdmx_no_comparable": no_comparable, "desviacion_justificada": justificada,
        "alerta": alerta, "estado": "revisada" if t.get("reviso") else "borrador",
        "elaboro": t.get("elaboro", ""), "reviso": t.get("reviso", ""), "fecha_base": c["par"]["fecha_base"],
        "supuestos": " | ".join(t.get("supuestos", [])),
    }


def filas_detalle(c):
    clave = c["def"]["clave_cb"]
    mats, n = [], 0
    for m in c["materiales"]:
        n += 1
        if m["tipo"] == "insumo":
            mats.append({"clave_cb": clave, "renglon": n, "tipo": "insumo", "clave": m["clave"], "basico": "",
                         "descripcion": m["descripcion"], "unidad": m["unidad"], "cantidad_base": m["cantidad_base"],
                         "desperdicio": m["desperdicio"], "cantidad": round(m["cantidad"], 6),
                         "precio": m["precio"], "importe": m["importe"], "nota": m["nota"]})
        else:
            mats.append({"clave_cb": clave, "renglon": n, "tipo": "basico", "clave": m["clave"], "basico": "",
                         "descripcion": m["descripcion"], "unidad": m["unidad"], "cantidad_base": m["cantidad_base"],
                         "desperdicio": m["desperdicio"], "cantidad": round(m["cantidad"], 6),
                         "precio": None, "importe": m["importe"], "nota": m["nota"]})
            for k in m["componentes"]:
                mats.append({"clave_cb": clave, "renglon": n, "tipo": "componente", "clave": k["clave"],
                             "basico": m["clave"], "descripcion": k["descripcion"], "unidad": k["unidad"],
                             "cantidad_base": k["coeficiente"], "desperdicio": 0, "cantidad": round(k["cantidad"], 6),
                             "precio": k["precio"], "importe": k["importe"], "nota": ""})
    mo = [{"clave_cb": clave, "categoria": q["categoria"], "nombre": q["nombre"], "personas": q["personas"],
           "salario_jornada": q["salario"], "rendimiento": c["rendimiento"], "fuente": c["fuente_rend"]}
          for q in c["cuadrilla"]]
    eq = [{"clave_cb": clave, "descripcion": q["descripcion"], "base": q["base"], "importe": q["importe"]}
          for q in c["equipo"]]
    return mats, mo, eq


def construir_base(claves=None):
    """Calcula todas las tarjetas de tarjetas/*.json y escribe la base de datos
    (SQLite y Excel). Devuelve (filas de tarjetas, errores)."""
    import sqlite3

    ref = cargar_referencias()
    claves = claves or sorted(p.stem for p in TARJETAS.glob("*.json"))
    tarjetas, materiales, mano_obra, equipo, errores = [], [], [], [], []
    for clave in claves:
        try:
            c = calcular(json.loads((TARJETAS / f"{clave}.json").read_text(encoding="utf-8")), ref)
        except (ErrorTarjeta, KeyError, ValueError) as err:
            errores.append((clave, str(err)))
            continue
        tarjetas.append(fila_tarjeta(c))
        m, o, e = filas_detalle(c)
        materiales += m
        mano_obra += o
        equipo += e
    tablas = {
        "tarjetas": tarjetas, "materiales": materiales, "mano_obra": mano_obra, "herramienta_equipo": equipo,
        "insumos": list(ref["insumos"].values()),
        "basicos": [{"clave": k, "descripcion": b["descripcion"], "unidad": b["unidad"], "insumo": i, "coeficiente": q}
                    for k, b in ref["basicos"].items() for i, q in b["componentes"]],
        "parametros": [{"parametro": k, "valor": v} for k, v in ref["parametros"].items() if k != "categorias"]
                      + [{"parametro": f"salario {k} ({n})", "valor": s}
                         for k, (n, s) in ref["parametros"]["categorias"].items()],
    }
    BASE_SQLITE.parent.mkdir(parents=True, exist_ok=True)
    if BASE_SQLITE.exists():
        BASE_SQLITE.unlink()
    con = sqlite3.connect(BASE_SQLITE)
    for nombre, filas in tablas.items():
        if not filas:
            continue
        cols = list(filas[0].keys())
        definicion = ", ".join('"' + x + '"' for x in cols)
        con.execute(f"CREATE TABLE {nombre} ({definicion})")
        con.executemany(f"INSERT INTO {nombre} VALUES ({', '.join('?' for _ in cols)})",
                        [[f.get(x) for x in cols] for f in filas])
    con.commit()
    con.close()
    _base_excel(tablas)
    return tarjetas, errores


def _base_excel(tablas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in tablas.items():
        ws = wb.create_sheet(nombre)
        if not filas:
            ws.append(["sin datos"])
            continue
        cols = list(filas[0].keys())
        ws.append(cols)
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
        for f in filas:
            ws.append([f.get(x) for x in cols])
        for i, col in enumerate(cols, start=1):
            ancho = 60 if col in ("descripcion", "supuestos", "nota", "fuente") else max(10, min(28, len(col) + 2))
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(BASE_EXCEL)


# --- Programa -------------------------------------------------------------------------

def documento(clave, pdf=True, excel=True):
    """PDF y/o Excel de una sola tarjeta, bajo pedido."""
    ref = cargar_referencias()
    c = calcular(json.loads((TARJETAS / f"{clave}.json").read_text(encoding="utf-8")), ref)
    SALIDA.mkdir(parents=True, exist_ok=True)
    if excel:
        excel_tarjeta(c, SALIDA / f"TPU-{clave}.xlsx")
    if pdf:
        pdf_desde_html(html_tarjeta(c), SALIDA / f"TPU-{clave}.pdf")
    return c


def validar(clave):
    """Calcula una tarjeta y muestra el resultado contra las referencias (lo usa el agente)."""
    ref = cargar_referencias()
    c = calcular(json.loads((TARJETAS / f"{clave}.json").read_text(encoding="utf-8")), ref)
    f = fila_tarjeta(c)
    print(f"{clave}  {f['descripcion'][:70]}")
    print(f"  GuBIM {f['gubim']} {f['gubim_nombre']} · unidad {f['unidad']}")
    print(f"  Materiales {f['materiales']:,.2f} | M.O. {f['mano_obra']:,.2f} | Herr./equipo {f['herramienta_equipo']:,.2f}"
          f" | Costo directo {f['costo_directo']:,.2f} | P.U. {f['pu']:,.2f}")
    for nombre, valor, dif in (("Construbase actualizado", f["pu_actualizado"], f["dif_vs_actualizado"]),
                               (f"Tabulador CDMX {f['clave_cdmx']}", f["pu_cdmx"], f["dif_vs_cdmx"])):
        if valor:
            print(f"  {nombre}: {valor:,.2f}  (diferencia {dif:+.1%})")
    print(f"  Rendimiento {f['rendimiento']} {f['unidad']}/jornada; implícito en el precio CDMX: "
          f"{f['rendimiento_implicito_cdmx'] or 'no aplica (sin par CDMX o precio no comparable)'}")
    if f["cdmx_no_comparable"]:
        print(f"  CDMX no comparable: {f['cdmx_no_comparable']}")
    if f["referencia_sustituta"]:
        print(f"  Referencia sustituta: {f['referencia_sustituta']}")
    if f["desviacion_justificada"]:
        print(f"  Desviación justificada: {f['desviacion_justificada']}")
    print(f"  {f['alerta'] or 'Sin alerta'}")
    return f


USO = """Uso:
  python3 scripts/tarjeta_pu.py                 reconstruye la base de datos con todas las tarjetas
  python3 scripts/tarjeta_pu.py --validar CLAVE  calcula una tarjeta y la compara con las referencias
  python3 scripts/tarjeta_pu.py --validar-lote CLAVES|archivo  resumen de muchas tarjetas
  python3 scripts/tarjeta_pu.py --pdf CLAVE      PDF y Excel de una tarjeta (bajo pedido)
  python3 scripts/tarjeta_pu.py --pendientes [texto]  conceptos sin tarjeta (filtra por capítulo o subcapítulo)"""


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        filas, errores = construir_base()
        print(f"Base de datos: {len(filas)} tarjetas -> {BASE_SQLITE.relative_to(RAIZ)} y {BASE_EXCEL.relative_to(RAIZ)}")
        for clave, err in errores:
            print(f"  ERROR {clave}: {err}")
        alertas = [f for f in filas if f["alerta"]]
        print(f"  {len(alertas)} con alerta de precio")
        distintas = [f["clave_cb"] for f in filas if f["gubim"] != f["gubim_catalogo"]]
        if distintas:
            print(f"  {len(distintas)} con GuBIM distinto del catálogo: {', '.join(distintas[:20])}")
    elif args[0] == "--validar":
        for clave in args[1:]:
            validar(clave)
    elif args[0] == "--validar-lote":
        # Resumen compacto de muchas tarjetas: una línea por tarjeta y las alertas al final.
        ref = cargar_referencias()
        claves = args[1:]
        if len(claves) == 1 and Path(claves[0]).exists():
            claves = Path(claves[0]).read_text().split()
        alertas, errores = [], []
        for clave in claves:
            ruta = TARJETAS / f"{clave}.json"
            if not ruta.exists():
                errores.append(f"{clave}: sin tarjeta")
                continue
            try:
                f = fila_tarjeta(calcular(json.loads(ruta.read_text(encoding="utf-8")), ref))
            except (ErrorTarjeta, KeyError, ValueError) as err:
                errores.append(f"{clave}: {err}")
                continue
            dif = f["dif_vs_referencia"]
            if f["gubim"] != f["gubim_catalogo"]:
                errores.append(f"{clave}: GuBIM {f['gubim']} distinto del catálogo ({f['gubim_catalogo']}); "
                               "regenera la tabla (python3 scripts/generar_tabla.py) o ajusta la regla")
            print(f"{clave} {f['unidad']:<4} PU {f['pu']:>11,.2f}  ref {f['referencia_validacion'] or '-':<16} "
                  f"{'' if dif is None else f'{dif:+.0%}':>6}  MO {f['mano_obra'] / f['costo_directo']:.0%}"
                  f"  {f['alerta'] or ('justificada' if f['desviacion_justificada'] else '')}")
            if f["alerta"]:
                alertas.append(clave)
        print(f"{len(claves)} tarjetas, {len(alertas)} con alerta, {len(errores)} con error")
        for e in errores:
            print("  ERROR", e)
    elif args[0] == "--pdf":
        for clave in args[1:]:
            documento(clave)
            print(f"salida/tarjetas/TPU-{clave}.pdf y .xlsx")
    elif args[0] == "--pendientes":
        ref = cargar_referencias()
        hechas = {p.stem for p in TARJETAS.glob("*.json")}
        filtro = " ".join(args[1:]).upper()
        pend = [r for r in ref["catalogo"].values() if r["clave_cb"] not in hechas and r["clave_gubim"]
                and (not filtro or filtro in r["capitulo"] or filtro in r["subcapitulo"])]
        print(f"{len(pend)} conceptos sin tarjeta")
        for r in pend[:40]:
            print(f"  {r['clave_cb']}  {r['unidad']:<4} {r['descripcion'][:90]}")
    else:
        sys.exit(USO)
