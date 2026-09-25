"""Genera la tabla Construbase -> GuBIMclass con precios y rendimientos.

Salidas (carpeta salida/):
  Construbase_GuBIM.xlsx  libro de trabajo con fórmulas editables
  construbase_gubim.csv   misma tabla en texto plano (precio 2017 y actualizado con los factores por defecto)

Uso:  python3 scripts/generar_tabla.py
"""
import collections
import csv
import math
import statistics
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import actualizacion
import comparar_tabulador
import leer_construbase
import reglas_gubim
import rendimientos
from reglas_gubim import normalizar

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "salida"
GUBIM = RAIZ / "fuentes" / "GuBIMclass_v1.2_ES.txt"

LETRA_SECCION = {"Edificación": "E", "Urbanización": "U", "Base Intelimat": "I"}
MO_ALERTA = 1.0     # mano de obra estimada > 100 % del precio
Z_ATIPICO = 3.5     # z robusto (log del precio) dentro de subcapítulo, unidad y actividad
VECES_ATIPICO = 3   # y además a más de 3 veces (o menos de 1/3) de la mediana

ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO_ENC = PatternFill("solid", fgColor="1F4E78")
FONDO_PARAM = PatternFill("solid", fgColor="FFF2CC")


def leer_gubim():
    gub = {}
    for linea in GUBIM.read_text(encoding="utf-8").splitlines():
        partes = linea.split("\t")
        if len(partes) >= 3:
            gub[partes[0]] = (partes[1], int(partes[2]))
    return gub


def asignar_claves(conceptos):
    """Clave propia: letra de sección + nº capítulo + nº subcapítulo + consecutivo."""
    caps, subs, cont = {}, {}, collections.Counter()
    for c in conceptos:
        kc = (c["seccion"], c["capitulo"])
        caps.setdefault(kc, len([k for k in caps if k[0] == c["seccion"]]) + 1)
        ks = kc + (c["subcapitulo"],)
        subs.setdefault(ks, len([k for k in subs if k[:2] == kc]) + 1)
        cont[ks] += 1
        c["clave"] = "{}{:02d}.{:02d}.{:04d}".format(
            LETRA_SECCION[c["seccion"]], caps[kc], subs[ks], cont[ks])


def detectar_alertas(conceptos):
    # Precios atípicos: se compara cada concepto con los de su mismo
    # subcapítulo, unidad y actividad (p. ej. excavaciones con excavaciones).
    grupos = collections.defaultdict(list)
    for c in conceptos:
        grupos[(c["seccion"], c["capitulo"], c["subcapitulo"], c["unidad"],
                (c["actividad"] or "").split("-")[0])].append(c)
    for grupo in grupos.values():
        if len(grupo) < 5:
            continue
        logs = [math.log(c["precio_2017"]) for c in grupo]
        med = statistics.median(logs)
        mad = statistics.median(abs(x - med) for x in logs)
        if mad == 0:
            continue
        for c, x in zip(grupo, logs):
            z = 0.6745 * (x - med) / mad
            veces = math.exp(x - med)
            if abs(z) > Z_ATIPICO and not 1 / VECES_ATIPICO < veces < VECES_ATIPICO:
                c["alertas"].append(
                    "Precio atípico: {:.1f}x la mediana de conceptos similares ({})".format(
                        veces, "alto" if z > 0 else "bajo"))
    # Descripciones repetidas.
    rep = collections.defaultdict(list)
    for c in conceptos:
        rep[(normalizar(c["descripcion"]), c["unidad"])].append(c)
    for grupo in rep.values():
        if len(grupo) < 2:
            continue
        precios = {c["precio_2017"] for c in grupo}
        otras = ", ".join(c["clave"] for c in grupo)
        for c in grupo:
            if len(precios) > 1:
                c["alertas"].append("Duplicado con precio distinto: " + otras)
            else:
                c["alertas"].append("Duplicado idéntico: " + otras)
    for c in conceptos:
        if c["unidad"] != c["unidad_origen"].upper():
            c["alertas"].append("Unidad normalizada ({} -> {})".format(c["unidad_origen"], c["unidad"]))
        if c["confianza"] == "baja":
            c["alertas"].append("Clave GuBIM genérica: revisar")


def construir():
    conceptos = leer_construbase.leer()
    gub = leer_gubim()
    asignar_claves(conceptos)
    for c in conceptos:
        c["gubim"], c["confianza"] = reglas_gubim.clasificar(c)
        c["actividad"], c["actividad_desc"], c["cuadrilla"], c["rendimiento"] = rendimientos.asignar(c)
        c["alertas"] = []
        costo = rendimientos.CUADRILLAS[c["cuadrilla"]][2] if c["cuadrilla"] else None
        c["mo_pct"] = costo / c["rendimiento"] / c["precio_2017"] if costo else None
        # Actualización por componentes con los factores por defecto (mismo
        # cálculo que las fórmulas del Excel).
        mo = min(c["mo_pct"] or 0, actualizacion.TOPE_MO)
        c["mo_act"] = round(c["precio_2017"] * mo * actualizacion.FACTOR_MO, 2)
        c["mat_act"] = round(c["precio_2017"] * (1 - mo) * actualizacion.FACTOR_MATERIALES, 2)
        c["pu_act"] = round(c["mo_act"] + c["mat_act"], 2)
    detectar_alertas(conceptos)
    return conceptos, gub


def encabezar(ws, columnas, anchos):
    ws.append(columnas)
    for i, ancho in enumerate(anchos, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font, celda.fill = ENCABEZADO, FONDO_ENC
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"


def hoja_leame(wb, n):
    ws = wb.active
    ws.title = "Léame"
    ws.column_dimensions["A"].width = 120
    lineas = [
        ("Construbase 2017 → GuBIMclass v1.2: precios y rendimientos", True),
        ("", False),
        ("Hojas", True),
        ("Parámetros: factores de actualización y costo por jornada de las cuadrillas (celdas amarillas).", False),
        ("Fuentes: de dónde salen los factores de actualización.", False),
        ("Tabulador CDMX: validación contra el tabulador de precios unitarios de la CDMX y ajustes por grupo.", False),
        (f"Catálogo: los {n} conceptos con clave propia, clave GuBIM, precio 2017, precio actualizado y rendimiento.", False),
        ("Rendimientos: tabla de actividades con cuadrilla y rendimiento (unidades por jornada de 8 h). Editable.", False),
        ("GuBIM: las 533 claves de GuBIMclass con el número de conceptos asignados a cada una.", False),
        ("Resumen: estadísticas de precio por capítulo y reparto de la confianza de la clave GuBIM.", False),
        ("Alertas: conceptos a revisar (precio atípico, duplicados, mano de obra estimada mayor al precio, clave genérica).", False),
        ("", False),
        ("Cómo se actualizan los precios (" + actualizacion.PERIODO + ", mercado Michoacán)", True),
        ("Por componentes: cada P.U. 2017 se divide en mano de obra (cuadrilla / rendimiento, con tope) y", False),
        ("materiales y equipo. La mano de obra se multiplica por el factor de M.O. y el resto por el de materiales.", False),
        ("Así funciona como una tarjeta paramétrica: sin matrices de Construbase, pero con la M.O. separada.", False),
        ("Si capturas en Parámetros el INPP Construcción residencial de Morelia (sep-2017 y mes actual) o un", False),
        ("factor manual, ese factor único sustituye a los dos factores por componente.", False),
        ("La columna 'P.U. con ajuste CDMX' corrige además los grupos donde el tabulador de la CDMX (jul-2026)", False),
        ("muestra una diferencia consistente (≥ 10 pares y más de ±10 %). Los ajustes son editables en esa hoja.", False),
        ("", False),
        ("Rendimientos: importante", True),
        ("El export de Construbase NO incluye las matrices de precios unitarios, por lo tanto no trae rendimientos.", False),
        ("Los rendimientos de esta tabla son valores de referencia de la práctica mexicana, asignados por tipo de", False),
        ("actividad y diámetro. Son un punto de partida para programar obra (4D) y deben validarse con las matrices", False),
        ("de Construbase o con datos propios de obra. La columna '% M.O. estimada' compara el costo de la cuadrilla", False),
        ("con el precio 2017: si supera 100 % el rendimiento o el precio es incongruente y se debe revisar.", False),
        ("", False),
        ("Confianza de la clave GuBIM", True),
        ("alta: el concepto identifica el elemento GuBIM. media: se deduce del contexto (p. ej. un muro puede ser", False),
        ("interior o de fachada). baja: clave genérica del capítulo. insumo: mezclas y concretos básicos que no se modelan.", False),
        ("", False),
        ("Uso con Revit", True),
        ("La clave GuBIM es el Assembly Code de Revit. Filtra el Catálogo por la clave del elemento para elegir el", False),
        ("concepto de precio; las horas-hombre por unidad sirven para estimar duraciones en la programación de obra.", False),
        ("", False),
        ("Generado con scripts/generar_tabla.py a partir de fuentes/construbase_2017.xlsx y fuentes/GuBIMclass_v1.2_ES.txt.", False),
    ]
    for texto, negrita in lineas:
        ws.append([texto])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=negrita, size=13 if negrita else 11)


def hoja_parametros(wb):
    ws = wb.create_sheet("Parámetros")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 22
    filas = [
        ("Actualización de precios (" + actualizacion.PERIODO + ")", None, None),
        ("Método A – índice único (si se captura, tiene prioridad)", None, None),
        ("INPP Construcción residencial Morelia, sep-2017", None, "INEGI, Índices de precios > Construcción residencial por ciudad"),
        ("INPP Construcción residencial Morelia, mes actual", None, "Misma serie y base que el valor de 2017"),
        ("Factor manual (opcional)", None, "Si se captura, sustituye al cálculo con INPP"),
        ("Factor único", '=IF(B6<>"",B6,IF(AND(N(B4)>0,N(B5)>0),B5/B4,""))',
         "Vacío = se usa el método B"),
        ("Método B – por componentes (se usa si el factor único está vacío)", None, None),
        ("Factor materiales y equipo", actualizacion.FACTOR_MATERIALES,
         "Estimación Michoacán; fuentes en la hoja Fuentes"),
        ("Factor mano de obra", actualizacion.FACTOR_MO,
         "Salario real de cuadrilla Michoacán 2026 / 2017"),
        ("Tope de % de mano de obra por concepto", actualizacion.TOPE_MO,
         "Limita la M.O. estimada cuando el rendimiento es incongruente"),
        ("Factor equivalente si % M.O. = 25 %", "=0.25*B10+0.75*B9", "Referencia para comparar con índices"),
    ]
    ws.append(["Parámetro", "Valor", "Nota"])
    for f in filas:
        ws.append(list(f))
    for r in (4, 5, 6, 9, 10, 11):
        ws.cell(row=r, column=2).fill = FONDO_PARAM
    for r in (2, 3, 8):
        ws.cell(row=r, column=1).font = Font(bold=True)
    for r in (7, 9, 10, 12):
        ws.cell(row=r, column=2).number_format = "0.000"
    ws["B11"].number_format = "0%"
    for c in ws[1]:
        c.font, c.fill = ENCABEZADO, FONDO_ENC

    ws.append([])
    ws.append(["Cuadrillas (costo por jornada, salario real con FSR)", None, None])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Clave", "Personas", "Descripción", "Costo jornada 2017", "Costo jornada actualizado"])
    fila_ini = ws.max_row + 1
    for clave, (desc, personas, costo) in rendimientos.CUADRILLAS.items():
        f = ws.max_row + 1
        ws.append([clave, personas, desc, costo, f'=IF(D{f}="","",ROUND(D{f}*IF($B$7<>"",$B$7,$B$10),2))'])
        ws.cell(row=f, column=4).fill = FONDO_PARAM
        ws.cell(row=f, column=4).number_format = "$#,##0.00"
        ws.cell(row=f, column=5).number_format = "$#,##0.00"
    celdas = {"unico": "Parámetros!$B$7", "mat": "Parámetros!$B$9",
              "mo": "Parámetros!$B$10", "tope": "Parámetros!$B$11"}
    return celdas, "Parámetros!$A${}:$D${}".format(fila_ini, ws.max_row)


def hoja_tabulador(wb, canasta, automaticos, ajustes, ruta_pdf):
    """Hoja con la validación contra el tabulador CDMX. Devuelve el rango de la
    tabla de ajustes por familia (columna 4 = ajuste aplicado, editable)."""
    ws = wb.create_sheet("Tabulador CDMX")
    anchos = [44, 12, 14, 60, 7, 12, 12, 12, 12, 11, 60, 12, 10, 40]
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a
    titulo = lambda t: (ws.append([t]), setattr(ws.cell(row=ws.max_row, column=1), "font", Font(bold=True, size=12)))

    titulo(f"Validación contra el Tabulador General de Precios Unitarios CDMX ({ruta_pdf.name})")
    ws.append(["P.U. de la CDMX sin cargos adicionales (3.627 %), con indirecto integrado 27.51 %, sin IVA. "
               "Cociente = P.U. CDMX / P.U. actualizado de esta tabla."])
    ws.append([])
    titulo("1. Ajuste por grupo (subcapítulo · familia); se aplica en la columna 'P.U. con ajuste CDMX' del Catálogo")
    ws.append(["Grupo", "Pares", "Mediana CDMX / actualizado", "Ajuste aplicado", "Criterio"])
    for c in ws[ws.max_row]:
        c.font, c.fill = ENCABEZADO, FONDO_ENC
    ini = ws.max_row + 1
    for fam, (n, med, sug) in sorted(ajustes.items(), key=lambda kv: -kv[1][0]):
        ws.append([fam, n, med, sug, "Ajuste sugerido" if sug != 1 else
                   ("Sin ajuste: diferencia menor a ±10 %" if n >= 10 else "Sin ajuste: menos de 10 pares")])
        ws.cell(row=ws.max_row, column=4).fill = FONDO_PARAM
        ws.cell(row=ws.max_row, column=3).number_format = "0.00"
        ws.cell(row=ws.max_row, column=4).number_format = "0.00"
    rango = f"'Tabulador CDMX'!$A${ini}:$D${ws.max_row}"

    def tabla(filas, extra):
        cab = ["Clave CB", "Capítulo", "Clave GuBIM", "Descripción Construbase", "Unidad", "P.U. 2017",
               "P.U. actualizado", "P.U. CDMX", "Cociente", "Clave CDMX", "Concepto CDMX"] + [e[0] for e in extra]
        ws.append(cab)
        for c in ws[ws.max_row]:
            c.font, c.fill = ENCABEZADO, FONDO_ENC
        for f in filas:
            ws.append([f["clave_cb"], f["capitulo"], f["clave_gubim"], f["descripcion_cb"], f["unidad"],
                       f["pu_2017"], f["pu_actualizado"], f["pu_tabulador"], f["cociente"],
                       f["clave_tabulador"], f["concepto_tabulador"]] + [f[e[1]] for e in extra])
            for col, fmt in ((6, "$#,##0.00"), (7, "$#,##0.00"), (8, "$#,##0.00"), (9, "0.00")):
                ws.cell(row=ws.max_row, column=col).number_format = fmt

    ws.append([])
    titulo(f"2. Canasta curada: {len(canasta)} conceptos típicos de vivienda emparejados a mano")
    tabla(sorted(canasta, key=lambda f: f["cociente"]), [("Equivalencia", "equivalencia"), ("Nota", "nota")])
    ws.append([])
    titulo(f"3. Pares automáticos con similitud ≥ 85: {len(automaticos)}")
    tabla(sorted(automaticos, key=lambda f: (f["capitulo"], f["cociente"])), [("Similitud", "similitud")])
    return rango


def hoja_fuentes(wb):
    ws = wb.create_sheet("Fuentes")
    encabezar(ws, ["Fuente", "Dato", "Uso en la actualización"], [48, 70, 70])
    for fila in actualizacion.FUENTES:
        ws.append(list(fila))
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    ws.append(["Pendiente", "Capturar el INPP Construcción residencial de Morelia (sep-2017 y último mes) en Parámetros",
               "Es el índice oficial más cercano al mercado de Michoacán; sustituye a la estimación"])


def hoja_rendimientos(wb):
    ws = wb.create_sheet("Rendimientos")
    encabezar(ws, ["Actividad", "Descripción", "Unidad", "Cuadrilla", "Rendimiento (u/jornada)",
                   "Jornadas por unidad", "Fuente"], [28, 55, 9, 10, 14, 12, 40])
    for (fam, u, tam), (desc, cuad, rend) in rendimientos.ACTIVIDADES.items():
        clave = "-".join(p for p in (fam, u, tam) if p)
        fila = ws.max_row + 1
        ws.append([clave, desc, u, cuad, rend, f"=1/E{fila}",
                   "Referencia típica, por validar con matrices u obra"])
        ws.cell(row=fila, column=5).fill = FONDO_PARAM
        ws.cell(row=fila, column=6).number_format = "0.0000"
    ws.auto_filter.ref = ws.dimensions
    return "Rendimientos!$A$2:$E${}".format(ws.max_row)


def hoja_catalogo(wb, conceptos, gub, celdas, rango_cuad, rango_rend, rango_ajustes=None):
    ws = wb.create_sheet("Catálogo")
    columnas = [
        # (clave interna, encabezado, ancho, formato)
        ("clave", "Clave CB", 13, None),
        ("seccion", "Sección", 12, None),
        ("capitulo", "Capítulo", 20, None),
        ("subcapitulo", "Subcapítulo", 22, None),
        ("descripcion", "Descripción", 70, None),
        ("unidad", "Unidad", 7, None),
        ("pu17", "P.U. 2017", 12, "$#,##0.00"),
        ("mo_act", "% M.O. para actualizar", 11, "0%"),
        ("mo_imp", "M.O. actualizada", 12, "$#,##0.00"),
        ("mat_imp", "Materiales y equipo actualizados", 13, "$#,##0.00"),
        ("pu_act", "P.U. actualizado", 13, "$#,##0.00"),
        ("ajuste", "Ajuste CDMX", 10, "0.00"),
        ("pu_cdmx", "P.U. con ajuste CDMX", 13, "$#,##0.00"),
        ("gubim", "Clave GuBIM", 12, None),
        ("gubim_desc", "Descripción GuBIM", 34, None),
        ("nivel", "Nivel", 6, None),
        ("confianza", "Confianza", 9, None),
        ("actividad", "Actividad", 22, None),
        ("grupo", "Grupo de ajuste CDMX", 30, None),
        ("cuadrilla", "Cuadrilla", 8, None),
        ("rend", "Rendimiento (u/jornada)", 12, None),
        ("jornadas", "Jornadas por unidad", 11, "0.0000"),
        ("hh", "Horas-hombre por unidad", 11, "0.00"),
        ("mo17", "% M.O. estimada 2017", 10, "0%"),
        ("alertas", "Alertas", 50, None),
        ("fila", "Fila origen", 8, None),
    ]
    if not rango_ajustes:
        columnas = [c for c in columnas if c[0] not in ("ajuste", "pu_cdmx")]
    L = {k: get_column_letter(i) for i, (k, *_ ) in enumerate(columnas, start=1)}
    encabezar(ws, [c[1] for c in columnas], [c[2] for c in columnas])
    unico, fmat, fmo, tope = celdas["unico"], celdas["mat"], celdas["mo"], celdas["tope"]
    for c in conceptos:
        f = ws.max_row + 1
        desc_g, nivel = gub.get(c["gubim"], ("", None)) if c["gubim"] else ("", None)
        act = c["actividad"] or ""
        col = lambda k: f"{L[k]}{f}"
        costo_cuad = f"VLOOKUP({col('cuadrilla')},{rango_cuad},4,0)"
        valores = {
            "clave": c["clave"], "seccion": c["seccion"], "capitulo": c["capitulo"],
            "subcapitulo": c["subcapitulo"], "descripcion": c["descripcion"], "unidad": c["unidad"],
            "pu17": c["precio_2017"],
            "mo_act": f'=IF(N({col("mo17")})=0,0,MIN({col("mo17")},{tope}))',
            "mo_imp": f'=ROUND({col("pu17")}*{col("mo_act")}*IF({unico}<>"",{unico},{fmo}),2)',
            "mat_imp": f'=ROUND({col("pu17")}*(1-{col("mo_act")})*IF({unico}<>"",{unico},{fmat}),2)',
            "pu_act": f'={col("mo_imp")}+{col("mat_imp")}',
            "ajuste": f'=IFERROR(VLOOKUP({col("grupo")},{rango_ajustes},4,0),1)' if rango_ajustes else "",
            "pu_cdmx": f'=ROUND({col("pu_act")}*{col("ajuste")},2)' if rango_ajustes else "",
            "grupo": comparar_tabulador.grupo(c),
            "gubim": c["gubim"] or "", "gubim_desc": desc_g, "nivel": nivel,
            "confianza": c["confianza"], "actividad": act, "cuadrilla": c["cuadrilla"] or "",
            "rend": f'=IFERROR(VLOOKUP({col("actividad")},{rango_rend},5,0),"")' if act else "",
            "jornadas": f'=IF({col("rend")}="","",1/{col("rend")})' if act else "",
            "hh": f'=IF({col("rend")}="","",VLOOKUP({col("cuadrilla")},{rango_cuad},2,0)*8/{col("rend")})' if act else "",
            "mo17": (f'=IF(OR({col("rend")}="",{costo_cuad}=""),"",'
                     f'{costo_cuad}/{col("rend")}/{col("pu17")})') if act else "",
            "alertas": "; ".join(c["alertas"]), "fila": c["fila_origen"],
        }
        ws.append([valores[k] for k, *_ in columnas])
        for i, (_, _, _, fmt) in enumerate(columnas, start=1):
            if fmt:
                ws.cell(row=f, column=i).number_format = fmt
    ultima = ws.max_row
    ws.auto_filter.ref = f"A1:{L['fila']}{ultima}"
    rojo = PatternFill("solid", fgColor="F8CBAD")
    ws.conditional_formatting.add(f"{L['mo17']}2:{L['mo17']}{ultima}",
                                  CellIsRule(operator="greaterThan", formula=["1"], fill=rojo))
    dv = DataValidation(type="list", formula1='"alta,media,baja,insumo"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{L['confianza']}2:{L['confianza']}{ultima}")


def hoja_gubim(wb, conceptos, gub):
    ws = wb.create_sheet("GuBIM")
    encabezar(ws, ["Clave GuBIM", "Descripción", "Nivel", "Conceptos asignados",
                   "P.U. 2017 mediano", "Unidad más frecuente"], [14, 60, 7, 12, 14, 12])
    por_clave = collections.defaultdict(list)
    for c in conceptos:
        if c["gubim"]:
            por_clave[c["gubim"]].append(c)
    for clave, (desc, nivel) in gub.items():
        lista = por_clave.get(clave, [])
        unidad = collections.Counter(c["unidad"] for c in lista).most_common(1)
        med = statistics.median(c["precio_2017"] for c in lista if c["unidad"] == unidad[0][0]) if lista else None
        ws.append([clave, "   " * (nivel - 1) + desc, nivel, len(lista), med, unidad[0][0] if unidad else ""])
        ws.cell(row=ws.max_row, column=5).number_format = "$#,##0.00"
        if nivel == 1:
            for cel in ws[ws.max_row]:
                cel.font = Font(bold=True)
    ws.auto_filter.ref = ws.dimensions


def hoja_resumen(wb, conceptos):
    ws = wb.create_sheet("Resumen")
    encabezar(ws, ["Sección", "Capítulo", "Conceptos", "P.U. mínimo", "P.U. mediano", "P.U. máximo",
                   "Confianza alta", "Confianza media", "Confianza baja", "Insumo",
                   "Alertas de precio", "M.O. > 100 %", "Factor de actualización medio"],
             [14, 34, 10, 12, 12, 14, 10, 10, 10, 8, 10, 10, 12])
    grupos = collections.defaultdict(list)
    for c in conceptos:
        grupos[(c["seccion"], c["capitulo"])].append(c)
    for (sec, cap), lista in grupos.items():
        precios = [c["precio_2017"] for c in lista]
        conf = collections.Counter(c["confianza"] for c in lista)
        ws.append([sec, cap, len(lista), min(precios), statistics.median(precios), max(precios),
                   conf["alta"], conf["media"], conf["baja"], conf["insumo"],
                   sum(any(a.startswith(("Precio atípico", "Duplicado con")) for a in c["alertas"]) for c in lista),
                   sum(1 for c in lista if c["mo_pct"] and c["mo_pct"] > MO_ALERTA),
                   sum(c["pu_act"] for c in lista) / sum(precios)])
        ws.cell(row=ws.max_row, column=13).number_format = "0.000"
        for col in (4, 5, 6):
            ws.cell(row=ws.max_row, column=col).number_format = "$#,##0.00"
    total = ws.max_row
    ws.append(["TOTAL", "", f"=SUM(C2:C{total})", "", "", "", f"=SUM(G2:G{total})", f"=SUM(H2:H{total})",
               f"=SUM(I2:I{total})", f"=SUM(J2:J{total})", f"=SUM(K2:K{total})", f"=SUM(L2:L{total})",
               sum(c["pu_act"] for c in conceptos) / sum(c["precio_2017"] for c in conceptos)])
    ws.cell(row=ws.max_row, column=13).number_format = "0.000"
    for cel in ws[ws.max_row]:
        cel.font = Font(bold=True)


def hoja_alertas(wb, conceptos):
    ws = wb.create_sheet("Alertas")
    encabezar(ws, ["Clave CB", "Tipo", "Detalle", "Capítulo", "Subcapítulo", "Descripción", "Unidad",
                   "P.U. 2017", "% M.O. estimada"], [13, 26, 50, 20, 22, 70, 7, 12, 10])
    for c in conceptos:
        filas = []
        for a in c["alertas"]:
            tipo = a.split(":")[0].split(" (")[0]
            if tipo in ("Unidad normalizada", "Duplicado idéntico"):
                continue
            filas.append((tipo, a))
        if c["mo_pct"] and c["mo_pct"] > MO_ALERTA:
            filas.append(("Rendimiento a revisar",
                          "M.O. estimada {:.0%} del precio con {} = {} {}/jornada".format(
                              c["mo_pct"], c["cuadrilla"], c["rendimiento"], c["unidad"])))
        for tipo, det in filas:
            ws.append([c["clave"], tipo, det, c["capitulo"], c["subcapitulo"], c["descripcion"],
                       c["unidad"], c["precio_2017"], c["mo_pct"]])
            ws.cell(row=ws.max_row, column=8).number_format = "$#,##0.00"
            ws.cell(row=ws.max_row, column=9).number_format = "0%"
    ws.auto_filter.ref = ws.dimensions


def escribir_csv(conceptos, gub):
    with open(SALIDA / "construbase_gubim.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["clave_cb", "seccion", "capitulo", "subcapitulo", "descripcion", "unidad",
                    "pu_2017", "clave_gubim", "descripcion_gubim", "confianza", "actividad",
                    "cuadrilla", "rendimiento_u_jornada", "horas_hombre_u", "mo_estimada_pct",
                    "mo_actualizada", "materiales_actualizados", "pu_actualizado", "pu_con_ajuste_cdmx", "alertas"])
        for c in conceptos:
            hh = (rendimientos.CUADRILLAS[c["cuadrilla"]][1] * 8 / c["rendimiento"]) if c["rendimiento"] else ""
            w.writerow([c["clave"], c["seccion"], c["capitulo"], c["subcapitulo"], c["descripcion"],
                        c["unidad"], c["precio_2017"], c["gubim"] or "",
                        gub.get(c["gubim"], ("",))[0] if c["gubim"] else "", c["confianza"],
                        c["actividad"] or "", c["cuadrilla"] or "", c["rendimiento"] or "",
                        round(hh, 4) if hh != "" else "",
                        round(c["mo_pct"], 3) if c["mo_pct"] else "",
                        c["mo_act"], c["mat_act"], c["pu_act"], c.get("pu_cdmx", c["pu_act"]),
                        "; ".join(c["alertas"])])


def main():
    SALIDA.mkdir(exist_ok=True)
    conceptos, gub = construir()
    wb = Workbook()
    hoja_leame(wb, len(conceptos))
    celdas, rango_cuad = hoja_parametros(wb)
    rango_rend = hoja_rendimientos(wb)
    rango_ajustes, ruta_tab = None, comparar_tabulador.tabulador_disponible()
    if ruta_tab:
        canasta, automaticos, ajustes = comparar_tabulador.evidencia(conceptos, ruta_tab)
        rango_ajustes = hoja_tabulador(wb, canasta, automaticos, ajustes, ruta_tab)
        for c in conceptos:
            aj = ajustes.get(comparar_tabulador.grupo(c), (0, 0, 1.0))[2]
            c["pu_cdmx"] = round(c["pu_act"] * aj, 2)
        with open(SALIDA / "pares_tabulador_cdmx.csv", "w", newline="", encoding="utf-8-sig") as f:
            campos = ["clave_cb", "clave_tabulador", "concepto_tabulador", "pu_tabulador", "cociente", "origen"]
            w = csv.DictWriter(f, fieldnames=campos, extrasaction="ignore")
            w.writeheader()
            for fila in canasta:
                w.writerow({**fila, "origen": "canasta " + fila["equivalencia"]})
            for fila in automaticos:
                w.writerow({**fila, "origen": f"automático (similitud {fila['similitud']})"})
        print(f"Tabulador CDMX: canasta {len(canasta)}, pares automáticos {len(automaticos)}, "
              f"ajustes {({k: v[2] for k, v in ajustes.items() if v[2] != 1})}")
    hoja_catalogo(wb, conceptos, gub, celdas, rango_cuad, rango_rend, rango_ajustes)
    hoja_gubim(wb, conceptos, gub)
    hoja_resumen(wb, conceptos)
    hoja_alertas(wb, conceptos)
    hoja_fuentes(wb)
    wb.move_sheet("Catálogo", offset=1 - wb.sheetnames.index("Catálogo"))
    wb.calculation.fullCalcOnLoad = True
    wb.save(SALIDA / "Construbase_GuBIM.xlsx")
    escribir_csv(conceptos, gub)
    conf = collections.Counter(c["confianza"] for c in conceptos)
    print(f"{len(conceptos)} conceptos; confianza {dict(conf)}")
    print("Claves GuBIM usadas:", len({c['gubim'] for c in conceptos if c['gubim']}))
    print("Con rendimiento:", sum(1 for c in conceptos if c["rendimiento"]))
    print("M.O. > 100 %:", sum(1 for c in conceptos if c["mo_pct"] and c["mo_pct"] > MO_ALERTA))


if __name__ == "__main__":
    main()
